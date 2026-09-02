from io import BytesIO

import docx
import openpyxl
from pypdf import PdfWriter

from app.services.ingestion import chunk_text, extract_text


def test_chunk_text_overlaps():
    chunks = chunk_text("word " * 500, size=100, overlap=20)
    assert len(chunks) > 1
    assert all(chunks)


def test_chunk_text_preserves_words_and_prefers_sentence_boundaries():
    text = "Admissions open in June. Scholarship applications follow in July. " * 8
    chunks = chunk_text(text, size=90, overlap=20)
    assert len(chunks) > 1
    assert all(len(chunk) <= 90 for chunk in chunks)
    assert all(not chunk.startswith("missions") for chunk in chunks)
    assert all("Admissio " not in chunk for chunk in chunks)
    assert all(chunk[-1].isalnum() or chunk[-1] in ".!?" for chunk in chunks)


def test_chunk_text_keeps_a_single_long_token_intact():
    token = "x" * 120
    assert chunk_text(token, size=50, overlap=10) == [token]


def test_text_markdown_and_csv_extraction():
    assert extract_text("notice.txt", b"Official notice") == "Official notice"
    assert extract_text("notice.md", b"# Official notice") == "# Official notice"
    assert extract_text("courses.csv", b"course,intake\nCSE,180\n") == "course | intake\nCSE | 180"


def test_docx_and_xlsx_extraction():
    doc_buffer = BytesIO()
    document = docx.Document()
    document.add_paragraph("Official admissions handbook")
    document.save(doc_buffer)
    assert "Official admissions handbook" in extract_text("handbook.docx", doc_buffer.getvalue())

    sheet_buffer = BytesIO()
    workbook = openpyxl.Workbook()
    workbook.active.append(["course", "intake"])
    workbook.active.append(["CSE", 180])
    workbook.save(sheet_buffer)
    extracted = extract_text("courses.xlsx", sheet_buffer.getvalue())
    assert "course | intake" in extracted and "CSE | 180" in extracted


def test_pdf_extraction_accepts_a_valid_pdf():
    buffer = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.write(buffer)
    assert extract_text("blank.pdf", buffer.getvalue()) == ""
