import csv
import io
import re
from pathlib import Path

import docx
import openpyxl
from pypdf import PdfReader


def extract_text(filename: str, content: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        reader = PdfReader(io.BytesIO(content))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if suffix == ".docx":
        document = docx.Document(io.BytesIO(content))
        return "\n".join(paragraph.text for paragraph in document.paragraphs)
    if suffix in {".xlsx", ".xlsm"}:
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        rows: list[str] = []
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                rows.append(" | ".join("" if cell is None else str(cell) for cell in row))
        return "\n".join(rows)
    if suffix == ".csv":
        decoded = content.decode("utf-8", errors="ignore")
        return "\n".join(" | ".join(row) for row in csv.reader(io.StringIO(decoded)))
    if suffix in {".txt", ".md"}:
        return content.decode("utf-8", errors="ignore")
    raise ValueError("Unsupported file type")


def chunk_text(text: str, size: int = 1200, overlap: int = 180) -> list[str]:
    """Create bounded, overlapping chunks without cutting through words.

    Sentence and paragraph boundaries are preferred. A single unusually long
    token is kept intact rather than silently corrupting searchable content.
    """
    normalized = re.sub(r"[ \t]+", " ", text.replace("\r\n", "\n").replace("\r", "\n")).strip()
    if not normalized or size <= 0:
        return []

    overlap = max(0, min(overlap, size - 1))
    segments = [
        segment.strip()
        for segment in re.split(r"(?<=[.!?])\s+|\n+", normalized)
        if segment.strip()
    ]
    chunks: list[str] = []

    current: list[str] = []
    current_length = 0

    def flush() -> None:
        nonlocal current, current_length
        if not current:
            return
        chunks.append(" ".join(current))
        carried: list[str] = []
        carried_length = 0
        for word in reversed(current):
            added = len(word) + (1 if carried else 0)
            if carried and carried_length + added > overlap:
                break
            if not carried and len(word) > overlap:
                break
            carried.insert(0, word)
            carried_length += added
        current = carried
        current_length = len(" ".join(current))

    for segment in segments:
        words = segment.split()
        for word in words:
            added = len(word) + (1 if current else 0)
            if current and current_length + added > size:
                flush()
                # Avoid an overlap-only chunk when the next token does not fit.
                if current and len(" ".join([*current, word])) > size:
                    current = []
                    current_length = 0
            current.append(word)
            current_length += len(word) + (1 if len(current) > 1 else 0)
        # Prefer ending a nearly-full chunk at the source sentence boundary.
        if current_length >= max(1, size - overlap):
            flush()

    if current and (not chunks or " ".join(current) != chunks[-1]):
        chunks.append(" ".join(current))
    return chunks
